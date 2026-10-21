#!/usr/bin/env python3
"""Inspect pinned official XLSX inputs, join immutable food keys, and generate a traceable numeric catalogue.
Raw files remain outside Git. No nutrient is filled with zero and no free-form source description is copied.
"""
import argparse
from decimal import Decimal, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
SOURCES={
 'ausnut-2023-food-details.xlsx':('https://www.foodstandards.gov.au/sites/default/files/2026-03/AUSNUT-2023-Food-details-4.xlsx','e7496d8fe2941ce80740f1731d7cb6e39836bc82d1a18d4007bc027780970293'),
 'ausnut-2023-food-nutrient-profiles.xlsx':('https://www.foodstandards.gov.au/sites/default/files/2025-09/AUSNUT%202023%20-%20Food%20nutrient%20profiles.xlsx?v=20250903','374650afc59951009d5ddf48fe35f712467458eb07eef513ea9f54c8f110f317'),
 'ausnut-2023-list-of-nutrients.xlsx':('https://www.foodstandards.gov.au/sites/default/files/2025-09/AUSNUT%202023%20-%20List%20of%20nutrients.xlsx?v=20250904','4d9c2547660afc7633ce1c7285612b8ed324d2be782493f9537f50565d564a9c')}
MAPPING={'energyKj100g':'Energy with dietary fibre (kJ)','proteinG100g':'Protein (g)','fatG100g':'Total fat (g)','carbsG100g':'Available carbohydrate, without sugar alcohols (g)','fiberG100g':'Dietary fibre (g)','sugarG100g':'Total sugars (g)','sodiumMg100g':'Sodium (Na) (mg)'}
POLICY='AUSNUT2023_ENERGY_WITH_FIBRE_CARB_EXCL_POLYOLS_CORE4_REQUIRED'
def digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()
def table(path,sheet,header_row,first_row):
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        ws=wb[sheet];header=list(next(ws.iter_rows(min_row=header_row,max_row=header_row,values_only=True)))
        return header,list(ws.iter_rows(min_row=first_row,values_only=True))
    finally:wb.close()
def number(value):
    if isinstance(value,bool) or not isinstance(value,(int,float,Decimal)):raise ValueError('Missing/non-numeric official nutrient; cannot substitute zero')
    n=Decimal(str(value))
    if not n.is_finite() or n<0 or n>Decimal('999999999.999'):raise ValueError('Invalid official nutrient range')
    n=n.quantize(Decimal('.001'),rounding=ROUND_HALF_UP)
    return int(n) if n==n.to_integral_value() else float(n)
def read_sources(directory,expected_hashes=None,expected_count=3741):
    directory=Path(directory);expected_hashes=expected_hashes or {n:v[1] for n,v in SOURCES.items()}
    for name,sha in expected_hashes.items():
        if digest(directory/name)!=sha:raise ValueError('Source SHA mismatch; inspect and assign a new revision: '+name)
    h,details=table(directory/'ausnut-2023-food-details.xlsx','Food details',3,5)
    for required in ['Survey ID','Public food key','Food name','Derivation']:
        if required not in h:raise ValueError('Missing exact food-details header: '+required)
    by_key={}
    for r in details:
        key=r[h.index('Public food key')]
        if not key or key in by_key:raise ValueError('Missing/duplicate public food key in source details')
        by_key[key]=(str(r[h.index('Survey ID')]),r[h.index('Food name')],r[h.index('Derivation')])
    h,rows=table(directory/'ausnut-2023-food-nutrient-profiles.xlsx','Food nutrient profiles',3,4)
    for required in ['Survey ID','Public food key','Food name',*MAPPING.values()]:
        if required not in h:raise ValueError('Missing exact nutrient header: '+required)
    nh,nutrients=table(directory/'ausnut-2023-list-of-nutrients.xlsx','TAB A - Nutrients - Foods ',3,4)
    units={r[0]:r[1] for r in nutrients}
    for name,unit in [('Energy with dietary fibre','kilojoules (kJ)'),('Protein','grams (g)'),('Total fat','grams (g)'),('Available carbohydrate, without sugar alcohols','grams (g)')]:
        if units.get(name)!=unit:raise ValueError('Nutrient definition/unit mismatch: '+name)
    result=[];keys=set();ids=set()
    for r in rows:
        key=r[h.index('Public food key')];survey=str(r[h.index('Survey ID')]);name=r[h.index('Food name')]
        if key in keys or survey in ids:raise ValueError('Duplicate immutable food key or survey ID')
        keys.add(key);ids.add(survey)
        if key not in by_key or by_key[key][:2]!=(survey,name):raise ValueError('Source join/name mismatch; inspect official revision')
        if not isinstance(name,str) or not name.strip() or len(name)>255:raise ValueError('Invalid official food name')
        result.append({'publicFoodKey':key,'surveyId':survey,'officialName':name,'displayName':name.strip(),'description':'Source derivation: '+str(by_key[key][2]),**{k:number(r[h.index(v)]) for k,v in MAPPING.items()}})
    if len(result)!=expected_count or len(by_key)!=expected_count or keys!=set(by_key):raise ValueError('Official food counts/key sets do not reconcile')
    return sorted(result,key=lambda r:r['publicFoodKey'])
def encode(value):return (json.dumps(value,ensure_ascii=False,sort_keys=True,indent=2)+'\n').encode()
def outputs(directory):
    rows=read_sources(directory);data=encode({'format':'fitnutri.food.v1','foods':rows});filename='ausnut-2023-foods-r1.json'
    manifest={'format':'fitnutri.food-manifest.v1','datasetName':'AUSNUT','release':'2023','revision':'2026-03-details_2025-09-nutrients-r1','sourceUrl':SOURCES['ausnut-2023-food-nutrient-profiles.xlsx'][0],'nutrientPolicy':POLICY,'recordsFile':filename,'recordsSha256':hashlib.sha256(data).hexdigest(),'expectedCount':len(rows),'sourceFiles':[{'filename':n,'url':u,'sha256':sha} for n,(u,sha) in SOURCES.items()], 'mapping':{'detailsSheet':'Food details','detailsHeaderRow':3,'detailsDataRow':5,'nutrientsSheet':'Food nutrient profiles','nutrientsHeaderRow':3,'nutrientsDataRow':4,'joinKeys':['Public food key','Survey ID'],'nutrientColumns':MAPPING,'nutrientDefinitionsSheet':'TAB A - Nutrients - Foods ','basis':'per 100 g edible portion','displayNameNormalisation':'Trim surrounding whitespace only; officialName retains the original source text.'},'sourceNote':'FSANZ AUSNUT 2023. Exact raw SHA and XLSX headers inspected. Energy includes fibre; available carbohydrate excludes sugar alcohols. Original food-name facts/nutrient values retained; free-form source descriptions and sampling prose are not copied. Attribution: © Food Standards Australia New Zealand. No endorsement is implied. Publisher copyright/licence page: https://www.foodstandards.gov.au/legal-policies/copyright.'}
    base=ROOT/'backend/src/main/resources/data'
    return {base/filename:data,base/'manifests/ausnut-2023-r1.json':encode(manifest)}
def main():
    parser=argparse.ArgumentParser();parser.add_argument('--source-dir',type=Path,default=ROOT/'.local/reference-sources');parser.add_argument('--write',action='store_true');args=parser.parse_args()
    for path,content in outputs(args.source_dir).items():
        if args.write:path.parent.mkdir(parents=True,exist_ok=True);path.write_bytes(content)
        elif not path.exists() or path.read_bytes()!=content:raise SystemExit('Official mapping drift: '+str(path))
    print(('Generated' if args.write else 'Verified')+' 3741 official AUSNUT foods, exact hashes/headers/key joins/core policy; no filled missing values.')
if __name__=='__main__':main()
