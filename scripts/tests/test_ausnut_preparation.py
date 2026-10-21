import hashlib
import importlib.util
from pathlib import Path
import tempfile
import unittest
from openpyxl import Workbook,load_workbook
spec=importlib.util.spec_from_file_location('prepare',Path(__file__).parents[1]/'data/prepare_ausnut.py');prepare=importlib.util.module_from_spec(spec);spec.loader.exec_module(prepare)
class PreparationTest(unittest.TestCase):
    def setUp(self):self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name);self.make()
    def tearDown(self):self.temp.cleanup()
    def workbook(self,name,sheet,header,row,start):
        wb=Workbook();ws=wb.active;ws.title=sheet
        for _ in range(2):ws.append([None])
        ws.append(header)
        while ws.max_row<start-1:ws.append([None])
        ws.append(row);wb.save(self.root/name)
    def make(self,core=0,header=None,detail_id=1):
        self.workbook('ausnut-2023-food-details.xlsx','Food details',['Survey ID','Public food key','Food name','Derivation'],[detail_id,'SYNTHETIC','Synthetic test food','Synthetic'],5)
        self.workbook('ausnut-2023-food-nutrient-profiles.xlsx','Food nutrient profiles',header or ['Survey ID','Public food key','Food name',*prepare.MAPPING.values()],[1,'SYNTHETIC','Synthetic test food',core,0,0,0,0,0,0],4)
        wb=Workbook();ws=wb.active;ws.title='TAB A - Nutrients - Foods ';ws.append([None]);ws.append([None]);ws.append(['Nutrient','Unit (per 100 g edible portion)'])
        for row in [('Energy with dietary fibre','kilojoules (kJ)'),('Protein','grams (g)'),('Total fat','grams (g)'),('Available carbohydrate, without sugar alcohols','grams (g)')]:ws.append(row)
        wb.save(self.root/'ausnut-2023-list-of-nutrients.xlsx')
    def read(self):return prepare.read_sources(self.root,{n:prepare.digest(self.root/n) for n in prepare.SOURCES},1)
    def test_real_zero_is_retained_and_identity_join_is_explicit(self):
        rows=self.read();self.assertEqual(1,len(rows));self.assertEqual('SYNTHETIC',rows[0]['publicFoodKey']);self.assertEqual('1',rows[0]['surveyId']);self.assertEqual(0,rows[0]['energyKj100g'])
    def test_missing_exact_header_is_rejected(self):
        self.make(header=['Survey ID','Public food key','Food name','invented energy column']);self.assertRaises(ValueError,self.read)
    def test_missing_nonnumeric_negative_core_is_not_substituted(self):
        for value in [None,'missing',-1]:
            with self.subTest(value=value):self.make(core=value);self.assertRaises(ValueError,self.read)
    def test_mismatched_source_identity_is_rejected(self):self.make(detail_id=2);self.assertRaises(ValueError,self.read)
    def test_duplicate_public_key_is_rejected_before_output(self):
        path=self.root/'ausnut-2023-food-nutrient-profiles.xlsx';wb=load_workbook(path);ws=wb['Food nutrient profiles'];ws.append([c.value for c in ws[4]]);wb.save(path);wb.close();self.assertRaises(ValueError,self.read)
    def test_source_fingerprint_mismatch_is_rejected(self):
        self.assertRaises(ValueError,prepare.read_sources,self.root,{n:'1'*64 for n in prepare.SOURCES},1)
if __name__=='__main__':unittest.main()
